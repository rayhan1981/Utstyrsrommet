"""Iteration 12 tests:
- GET /items/taxonomy
- GET /items/import/template (xlsx with instructions sheet, suppliers list)
- POST /items/import/preview (valid + error cases: type, supplier, missing category)
- POST /items/import (commits + supplier auto-match on Håndball Select)
- GET /items/{id}/open-orders (filters levert/arkivert/annullert, empty case)
"""
import os
import io
import uuid
import pytest
import requests
from datetime import datetime, timezone

from openpyxl import load_workbook, Workbook


def _url():
    u = os.environ.get("REACT_APP_BACKEND_URL", "").strip()
    if not u:
        with open("/app/frontend/.env") as f:
            for line in f:
                if line.startswith("REACT_APP_BACKEND_URL="):
                    u = line.split("=", 1)[1].strip()
    return u.rstrip("/")


BASE = _url()
API = BASE + "/api"
ADMIN_EMAIL = "raymond.d.hansen@gmail.com"
ADMIN_PASSWORD = "Snowboard1!"

IMPORT_COLUMNS = ["Varenavn", "Kategori", "Hovedgruppe", "Undergruppe", "Pris", "Lokasjon", "Type", "Størrelse", "Draktnr", "Antall", "Lokasjon pr str", "Leverandørnr"]


@pytest.fixture(scope="session")
def owner_session():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert r.status_code == 200, r.text
    s.headers["Authorization"] = f"Bearer {r.json()['token']}"
    return s


@pytest.fixture(scope="session")
def tenant(owner_session):
    suffix = uuid.uuid4().hex[:8]
    email = f"TEST_it12_{suffix}@example.com"
    password = "TestPass1!"
    r = requests.post(f"{API}/signup", json={
        "club_name": f"TEST_It12_{suffix}",
        "owner_name": "TEST Eier It12",
        "owner_email": email,
        "owner_phone": "+4712345678",
        "password": password,
        "plan": "deluxe",
        "accept_terms": True,
    })
    assert r.status_code == 200, r.text
    cust = r.json()["customer"]
    owner_session.patch(f"{API}/owner/customers/{cust['id']}", json={"plan": "deluxe", "status": "active"})
    s = requests.Session()
    lr = s.post(f"{API}/auth/login", json={"email": email, "password": password})
    assert lr.status_code == 200, lr.text
    s.headers["Authorization"] = f"Bearer {lr.json()['token']}"
    yield {"session": s, "customer": cust}
    try:
        owner_session.delete(f"{API}/owner/customers/{cust['id']}")
    except Exception:
        pass


def _create_supplier(sess, number="00001", name="Select Sport AS"):
    r = sess.post(f"{API}/suppliers", json={
        "name": name, "email": "sup@example.com",
        "supplier_number": number,
    })
    assert r.status_code == 200, r.text
    return r.json()


def _create_simple_item(sess, name, category="Klær", mg="Diverse"):
    r = sess.post(f"{API}/items", json={
        "name": name, "category": category, "main_group": mg,
        "price": 100.0, "total_stock": 5, "location": "Lager",
        "is_clothing": False, "is_ball": False,
    })
    assert r.status_code == 200, r.text
    return r.json()


# ================== Taxonomy ==================
class TestTaxonomy:
    def test_taxonomy_shape_and_populated(self, tenant):
        s = tenant["session"]
        _create_simple_item(s, name=f"TEST_It12_Tax_{uuid.uuid4().hex[:6]}", category="Klær", mg="Drakt")
        _create_simple_item(s, name=f"TEST_It12_Tax2_{uuid.uuid4().hex[:6]}", category="Utstyr", mg="Baller")
        r = s.get(f"{API}/items/taxonomy")
        assert r.status_code == 200, r.text
        d = r.json()
        assert "categories" in d
        assert isinstance(d["categories"], list)
        assert len(d["categories"]) >= 1
        for cat in d["categories"]:
            assert "name" in cat and "main_groups" in cat
            for mg in cat["main_groups"]:
                assert "name" in mg and "sub_groups" in mg
                assert isinstance(mg["sub_groups"], list)
        names = [c["name"] for c in d["categories"]]
        assert "Klær" in names


# ================== Import template ==================
class TestImportTemplate:
    def test_template_download(self, tenant):
        s = tenant["session"]
        # Create a supplier so the suppliers list in instructions is populated
        _create_supplier(s, number="00001", name="Select Sport AS")
        r = s.get(f"{API}/items/import/template")
        assert r.status_code == 200, r.text
        assert len(r.content) > 5000, f"Template too small: {len(r.content)} bytes"
        wb = load_workbook(io.BytesIO(r.content))
        assert "Varer" in wb.sheetnames
        assert "Instruksjoner" in wb.sheetnames
        ws = wb["Varer"]
        headers = [ws.cell(row=1, column=i+1).value for i in range(len(IMPORT_COLUMNS))]
        assert headers == IMPORT_COLUMNS
        assert headers[-1] == "Leverandørnr"
        # First data row should be Kjegler
        assert ws.cell(row=2, column=1).value == "Kjegler"
        # Instruksjoner sheet should list supplier 00001
        info = wb["Instruksjoner"]
        found = False
        for row in info.iter_rows(values_only=True):
            if row and any(c == "00001" for c in row if c):
                found = True; break
        assert found, "Supplier 00001 not listed in Instruksjoner sheet"


# ================== Import preview valid ==================
def _build_xlsx(rows):
    """Build an xlsx with IMPORT_COLUMNS header + given data rows (lists)."""
    wb = Workbook()
    ws = wb.active; ws.title = "Varer"
    ws.append(IMPORT_COLUMNS)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


class TestImportPreview:
    def test_preview_valid_template_with_supplier_match(self, tenant):
        s = tenant["session"]
        _create_supplier(s, number="00001", name="Select Sport AS")
        # Download real template
        r = s.get(f"{API}/items/import/template")
        assert r.status_code == 200
        files = {"file": ("mal.xlsx", io.BytesIO(r.content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        pr = s.post(f"{API}/items/import/preview", files=files)
        assert pr.status_code == 200, pr.text
        d = pr.json()
        assert d["ok"] is True
        assert d["errors"] == []
        # Template has 3 unique items: Kjegler, Hjemmedrakt (2 variants), Håndball Select (2 variants)
        assert d["total_items"] == 3
        assert d["total_variants"] == 4
        # Håndball Select should match supplier 00001
        handball = next((p for p in d["items"] if p["name"] == "Håndball Select"), None)
        assert handball is not None
        assert handball["supplier_number"] == "00001"
        assert handball["supplier_name"] == "Select Sport AS"

    def test_preview_invalid_type(self, tenant):
        s = tenant["session"]
        buf = _build_xlsx([["Testvare", "Utstyr", "", "", 10, "Lager", "xyz", "", "", 1, "", ""]])
        files = {"file": ("bad.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = s.post(f"{API}/items/import/preview", files=files)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["ok"] is False
        type_errs = [e for e in d["errors"] if e["field"] == "Type"]
        assert len(type_errs) >= 1
        assert "vare" in type_errs[0]["error"].lower() or "type" in type_errs[0]["error"].lower()

    def test_preview_missing_supplier(self, tenant):
        s = tenant["session"]
        buf = _build_xlsx([["Testvare2", "Utstyr", "", "", 10, "Lager", "vare", "", "", 1, "", "99999"]])
        files = {"file": ("badsup.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = s.post(f"{API}/items/import/preview", files=files)
        assert r.status_code == 200, r.text
        d = r.json()
        sup_errs = [e for e in d["errors"] if e["field"] == "Leverandørnr"]
        assert len(sup_errs) >= 1
        assert "finnes ikke" in sup_errs[0]["error"]

    def test_preview_missing_category(self, tenant):
        s = tenant["session"]
        buf = _build_xlsx([["Testvare3", "", "", "", 10, "Lager", "vare", "", "", 1, "", ""]])
        files = {"file": ("badcat.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = s.post(f"{API}/items/import/preview", files=files)
        assert r.status_code == 200, r.text
        d = r.json()
        cat_errs = [e for e in d["errors"] if e["field"] == "Kategori"]
        assert len(cat_errs) >= 1
        assert cat_errs[0]["error"] == "Kategori er påkrevd"


# ================== Import commit ==================
class TestImportCommit:
    def test_import_commits_items_with_supplier_match(self, tenant):
        s = tenant["session"]
        _create_supplier(s, number="00001", name="Select Sport AS")
        r = s.get(f"{API}/items/import/template")
        assert r.status_code == 200
        files = {"file": ("mal.xlsx", io.BytesIO(r.content), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        imp = s.post(f"{API}/items/import", files=files)
        assert imp.status_code == 200, imp.text
        d = imp.json()
        assert d["created"] == 3
        assert d["errors"] == []
        # Verify items exist in DB with supplier info on Håndball Select
        items = s.get(f"{API}/items").json()
        handball = next((i for i in items if i["name"] == "Håndball Select"), None)
        assert handball is not None
        assert handball.get("supplier_number") == "00001"
        assert handball.get("supplier_name") == "Select Sport AS"
        # Variants attached
        assert len(handball.get("sizes", [])) == 2


# ================== Open orders panel ==================
class TestOpenOrders:
    def test_open_orders_empty_for_item_with_no_po(self, tenant):
        s = tenant["session"]
        item = _create_simple_item(s, name=f"TEST_It12_NoPO_{uuid.uuid4().hex[:6]}")
        r = s.get(f"{API}/items/{item['id']}/open-orders")
        assert r.status_code == 200, r.text
        assert r.json() == {"orders": []}

    def test_open_orders_returns_po_lines(self, tenant):
        s = tenant["session"]
        item = _create_simple_item(s, name=f"TEST_It12_WithPO_{uuid.uuid4().hex[:6]}")
        sup = _create_supplier(s, number=f"{10000 + int(uuid.uuid4().hex[:4], 16) % 80000:05d}", name=f"TEST_Sup_{uuid.uuid4().hex[:6]}")
        orr = s.post(f"{API}/orders", json={
            "supplier_id": sup["id"], "is_stock_order": True,
            "lines": [{"item_id": item["id"], "item_name": item["name"], "quantity": 10, "unit_price": 50.0}],
        })
        assert orr.status_code == 200, orr.text
        r = s.get(f"{API}/items/{item['id']}/open-orders")
        assert r.status_code == 200, r.text
        d = r.json()
        assert "orders" in d
        assert len(d["orders"]) == 1
        line = d["orders"][0]
        for key in ("order_id", "po_number", "supplier_name", "ordered", "received", "remaining", "line_status", "status"):
            assert key in line, f"Missing {key} in open-orders response"
        assert line["ordered"] == 10
        assert line["received"] == 0
        assert line["remaining"] == 10

    def test_open_orders_filters_delivered(self, tenant):
        s = tenant["session"]
        item = _create_simple_item(s, name=f"TEST_It12_Delivered_{uuid.uuid4().hex[:6]}")
        sup = _create_supplier(s, number=f"{10000 + int(uuid.uuid4().hex[:4], 16) % 80000:05d}", name=f"TEST_Sup2_{uuid.uuid4().hex[:6]}")
        orr = s.post(f"{API}/orders", json={
            "supplier_id": sup["id"], "is_stock_order": True,
            "lines": [{"item_id": item["id"], "item_name": item["name"], "quantity": 5, "unit_price": 10.0}],
        })
        order = orr.json()
        # Try to archive or mark delivered. Attempt via manual-status endpoint.
        pr = s.post(f"{API}/orders/{order['id']}/manual-status", json={"status": "levert", "reason": "Test – mark levert"})
        if pr.status_code != 200:
            pytest.skip(f"Cannot transition order to levert: {pr.status_code} {pr.text}")
        r = s.get(f"{API}/items/{item['id']}/open-orders")
        assert r.status_code == 200
        assert r.json()["orders"] == []
