"""
Router for bestillinger, behovsskjema, leverandører, varemottak, behov og reservasjoner.

Inneholder:
- Send utstyrsliste (PDF) til lagleder
- Varemottak (stock-receipts: resolve-qr, list, create, pdf)
- Team needs (innmeldt behov: seasons, template, list, import, confirm, delete)
- Leverandører (suppliers: list, trash, restore, create, update, toggle-active, delete, orders)
- Behovsskjema (create, list, update, send, delete, import, lines, confirm, unlock, udekket)
- Purchase orders (create, list, by-po, receive, manual-status, resume-auto-status, pdf,
  notify-lagleder, send)
- Stock reservations
- Purchase history per lag

Hjelpefunksjoner som kun brukes her: _current_seasons, _build_needs_template_bytes,
_compute_need_coverage, _next_supplier_number, _next_po_number, _build_order_pdf, _send_sms.
Felles hjelpefunksjoner importeres via _shared.
"""
import io
import re
import base64 as _b64mod
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Dict, Any

from fastapi import Depends, HTTPException, Request, Response, UploadFile, File
from pydantic import BaseModel, EmailStr
from openpyxl import Workbook, load_workbook

from ._shared import (
    api, db, require_tenant, require_perm, require_admin, get_current_user,
    _log_activity, now_iso, new_id, _compute_full_number, send_tenant_email,
    _generate_team_equipment_pdf_for_tenant, logger,
    EMAIL_KEY,
    TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN, TWILIO_FROM_NUMBER,
)


# ============== SEND UTSTYRSLISTE (PDF) TIL LAGLEDER ==============
class TeamSendEquipmentIn(BaseModel):
    to: Optional[EmailStr] = None  # override; defaults to lagleder-epost
    cc: Optional[EmailStr] = None
    message: Optional[str] = None


@api.post("/teams/{team_id}/send-equipment-list")
async def send_team_equipment_list(team_id: str, data: TeamSendEquipmentIn, request: Request,
                                     user=Depends(require_perm("teams", "write"))):
    """Send utstyrsoversikt (PDF) rett til lagleder på epost."""
    tid = await require_tenant(request, user)
    pdf_bytes, team, club_name, _ = await _generate_team_equipment_pdf_for_tenant(tid, team_id, user)
    if not team:
        raise HTTPException(404, "Lag ikke funnet")
    tm = team.get("team_manager") or {}
    hc = team.get("head_coach") or {}
    recipient = {"name": tm.get("name") or hc.get("name") or team.get("contact_person") or "",
                 "email": tm.get("email") or hc.get("email") or team.get("email") or ""}
    to_email = (data.to or recipient["email"] or "").strip()
    if not to_email or not re.match(r"^[^@]+@[^@]+\.[^@]+$", to_email):
        raise HTTPException(400, "Mangler gyldig e-post til lagleder. Sett e-post på laget først, eller angi mottaker.")
    pdf_b64 = _b64mod.b64encode(pdf_bytes).decode()
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", team.get("team_name", "lag"))
    filename = f"utstyr_{safe}.pdf"
    msg = (data.message or "").strip()
    html = (
        f'<p>Hei {recipient["name"] or ""},</p>'
        f'<p>Vedlagt er oversikt over utstyr og eventuelt utestående beløp for <strong>{team.get("team_name")}</strong>.</p>'
        + (f'<p>{msg}</p>' if msg else '')
        + f'<p>Vennligst sjekk listen og returner utstyr før fristdato.</p>'
        f'<p>Mvh<br/>{club_name}</p>'
    )
    cc_list = []
    if data.cc: cc_list.append(data.cc)
    if user.get("email") and user["email"] != to_email and user["email"] not in cc_list:
        cc_list.append(user["email"])
    attachments = [{"filename": filename, "content": pdf_b64, "type": "application/pdf"}]
    res = await send_tenant_email(tid, to_email,
                                   f"Utstyrsoversikt – {team.get('team_name')}",
                                   html, attachments=attachments, cc=cc_list or None,
                                   sender_user=user)
    if not res.get("ok"):
        raise HTTPException(502, f"E-post feilet: {res.get('error')}")
    await db.teams.update_one({"id": team_id, "tenant_id": tid}, {"$set": {
        "equipment_list_sent_at": now_iso(), "equipment_list_sent_to": to_email,
    }})
    return {"ok": True, "sent_to": to_email, "provider_id": res.get("provider_id"), "method": res.get("method")}


# ============== VAREMOTTAK (Goods Receipt) ==============
class StockReceiptIn(BaseModel):
    request_id: Optional[str] = None
    item_id: str
    size: Optional[str] = None
    jersey_number: Optional[str] = None
    quantity: int
    order_ref: str = ""
    comment: str = ""


class QrResolveIn(BaseModel):
    code: str


@api.post("/stock-receipts/resolve-qr")
async def resolve_receipt_qr(data: QrResolveIn, request: Request, user=Depends(require_perm("inventory", "write"))):
    tid = await require_tenant(request, user)
    code = (data.code or "").strip()
    if not code:
        raise HTTPException(400, "Mangler QR-kode")
    item_id = code; size = None
    if "#" in code:
        item_id, size = code.split("#", 1)
    item = await db.items.find_one({"id": item_id, "tenant_id": tid}, {"_id": 0})
    if not item:
        raise HTTPException(404, "QR-koden tilhører ikke en vare i systemet")
    return {"item": item, "size": size}


@api.get("/stock-receipts")
async def list_receipts(request: Request, user=Depends(require_perm("inventory", "read"))):
    tid = await require_tenant(request, user)
    receipts = await db.stock_receipts.find({"tenant_id": tid}, {"_id": 0}).sort("created_at", -1).limit(200).to_list(200)
    # Backfill full_number on-the-fly for gamle rader der feltet mangler.
    missing_item_ids = list({r.get("item_id") for r in receipts if not r.get("full_number") and r.get("item_id")})
    if missing_item_ids:
        items = await db.items.find({"id": {"$in": missing_item_ids}, "tenant_id": tid},
                                     {"_id": 0, "id": 1, "main_item_number": 1, "full_number": 1, "sizes": 1}).to_list(5000)
        by_id = {it["id"]: it for it in items}
        for r in receipts:
            if r.get("full_number"): continue
            it = by_id.get(r.get("item_id"))
            if not it: continue
            main = it.get("main_item_number") or it.get("full_number") or ""
            if r.get("size"):
                sz_entry = next((s for s in it.get("sizes") or [] if s.get("size") == r.get("size")), None)
                r["full_number"] = (sz_entry or {}).get("full_number") or _compute_full_number(main, r.get("size"), r.get("jersey_number"))
            else:
                r["full_number"] = main
    return receipts


@api.post("/stock-receipts")
async def create_receipt(data: StockReceiptIn, request: Request, user=Depends(require_perm("inventory", "write"))):
    tid = await require_tenant(request, user)
    if not isinstance(data.quantity, int) or data.quantity <= 0:
        raise HTTPException(400, "Mottatt antall må være et positivt heltall")
    if data.request_id:
        existing = await db.stock_receipts.find_one({"tenant_id": tid, "request_id": data.request_id}, {"_id": 0})
        if existing:
            return existing
    item = await db.items.find_one({"id": data.item_id, "tenant_id": tid})
    if not item:
        raise HTTPException(404, "Vare finnes ikke")
    qty_before = 0; qty_after = 0
    full_number = ""
    main_num = item.get("main_item_number") or item.get("full_number") or ""
    if item.get("is_clothing") or item.get("is_ball"):
        if not data.size:
            raise HTTPException(400, "Størrelse må oppgis for klær og baller")
        sizes = item.get("sizes") or []
        found = None
        for s in sizes:
            if s.get("size") == data.size and (
                not data.jersey_number or data.jersey_number in ("X", s.get("jersey_number"))
            ):
                found = s; break
        if not found:
            raise HTTPException(404, "Fant ikke angitt variant på varen")
        qty_before = int(found.get("quantity") or 0)
        found["quantity"] = qty_before + data.quantity
        qty_after = found["quantity"]
        new_total = sum(int(s.get("quantity") or 0) for s in sizes)
        full_number = found.get("full_number") or _compute_full_number(main_num, data.size, found.get("jersey_number"))
        await db.items.update_one({"id": data.item_id, "tenant_id": tid},
                                  {"$set": {"sizes": sizes, "total_stock": new_total}})
    else:
        qty_before = int(item.get("total_stock") or 0)
        qty_after = qty_before + data.quantity
        full_number = main_num
        await db.items.update_one({"id": data.item_id, "tenant_id": tid},
                                  {"$set": {"total_stock": qty_after}})
    receipt = {
        "id": new_id(), "tenant_id": tid, "request_id": data.request_id,
        "item_id": data.item_id, "item_name": item.get("name"),
        "full_number": full_number,
        "size": data.size, "jersey_number": data.jersey_number,
        "quantity": data.quantity, "qty_before": qty_before, "qty_after": qty_after,
        "order_ref": data.order_ref, "comment": data.comment,
        "user_id": user.get("id"), "user_name": user.get("name"),
        "created_at": now_iso(),
    }
    await db.stock_receipts.insert_one(receipt)
    receipt.pop("_id", None)
    return receipt


@api.get("/stock-receipts/pdf")
async def stock_receipts_pdf(
    request: Request,
    date_from: Optional[str] = None,   # ISO (YYYY-MM-DD)
    date_to: Optional[str] = None,     # ISO (YYYY-MM-DD)
    user=Depends(require_perm("inventory", "read")),
):
    """Varemottak-PDF: liste over mottatte varer i periode. Varenr som første kolonne."""
    tid = await require_tenant(request, user)
    q: Dict[str, Any] = {"tenant_id": tid}
    # Datofilter inkluderer hele to-dagen (<= 23:59:59)
    if date_from or date_to:
        rng: Dict[str, str] = {}
        if date_from: rng["$gte"] = f"{date_from}T00:00:00+00:00"
        if date_to:   rng["$lte"] = f"{date_to}T23:59:59+00:00"
        q["created_at"] = rng
    receipts = await db.stock_receipts.find(q, {"_id": 0}).sort("created_at", -1).to_list(5000)

    # Backfill full_number
    missing_ids = list({r.get("item_id") for r in receipts if not r.get("full_number") and r.get("item_id")})
    if missing_ids:
        items = await db.items.find({"id": {"$in": missing_ids}, "tenant_id": tid},
                                     {"_id": 0, "id": 1, "main_item_number": 1, "full_number": 1, "sizes": 1}).to_list(5000)
        by_id = {it["id"]: it for it in items}
        for r in receipts:
            if r.get("full_number"): continue
            it = by_id.get(r.get("item_id"))
            if not it: continue
            main = it.get("main_item_number") or it.get("full_number") or ""
            if r.get("size"):
                sz = next((s for s in it.get("sizes") or [] if s.get("size") == r.get("size")), None)
                r["full_number"] = (sz or {}).get("full_number") or _compute_full_number(main, r.get("size"), r.get("jersey_number"))
            else:
                r["full_number"] = main

    cust = await db.customers.find_one({"id": tid}, {"_id": 0}) or {}
    club = await db.clubs.find_one({"tenant_id": tid}, {"_id": 0}) or {}
    club_name = club.get("name") or cust.get("club_name") or "Klubben"

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas

    buf = io.BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    LEFT = 15 * mm; RIGHT = w - 15 * mm
    y = h - 20 * mm

    c.setFont("Helvetica-Bold", 18)
    c.drawString(LEFT, y, "VAREMOTTAK")
    c.setFont("Helvetica", 9)
    c.drawRightString(RIGHT, y, datetime.now(timezone.utc).strftime("%d.%m.%Y %H:%M"))
    y -= 6 * mm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(LEFT, y, club_name)
    y -= 5 * mm
    period = ""
    if date_from and date_to: period = f"Periode: {date_from} – {date_to}"
    elif date_from: period = f"Fra: {date_from}"
    elif date_to: period = f"Til: {date_to}"
    else: period = "Alle registrerte mottak"
    c.setFont("Helvetica", 9)
    c.drawString(LEFT, y, period)
    y -= 4 * mm
    c.drawString(LEFT, y, f"Antall linjer: {len(receipts)}")
    y -= 6 * mm
    c.line(LEFT, y, RIGHT, y); y -= 5 * mm

    # Kolonner (Varenr først): Varenr | Dato | Vare | Str/Nr | Antall | Før→Etter | Bestilling | Bruker
    col_varenr = LEFT
    col_dato = LEFT + 28 * mm
    col_vare = LEFT + 50 * mm
    col_str = LEFT + 100 * mm
    col_qty_r = LEFT + 122 * mm
    col_delta_r = LEFT + 145 * mm
    col_order = LEFT + 150 * mm
    col_user_r = RIGHT

    def _hdr():
        c.setFont("Helvetica-Bold", 8)
        c.drawString(col_varenr, y, "Varenr")
        c.drawString(col_dato, y, "Dato")
        c.drawString(col_vare, y, "Vare")
        c.drawString(col_str, y, "Str/Nr")
        c.drawRightString(col_qty_r, y, "Antall")
        c.drawRightString(col_delta_r, y, "Før→Etter")
        c.drawString(col_order, y, "Bestilling")
        c.drawRightString(col_user_r, y, "Bruker")
    _hdr()
    y -= 2 * mm; c.line(LEFT, y, RIGHT, y); y -= 4 * mm

    c.setFont("Helvetica", 8)
    total_qty = 0
    if not receipts:
        c.setFont("Helvetica-Oblique", 9)
        c.drawString(LEFT, y, "Ingen varemottak i valgt periode.")
    else:
        for r in receipts:
            if y < 25 * mm:
                c.showPage()
                y = h - 20 * mm
                _hdr()
                y -= 2 * mm; c.line(LEFT, y, RIGHT, y); y -= 4 * mm
                c.setFont("Helvetica", 8)
            try:
                dato = datetime.fromisoformat(r.get("created_at", "").replace("Z", "+00:00")).strftime("%d.%m.%Y")
            except Exception:
                dato = (r.get("created_at") or "")[:10]
            size_bits = []
            if r.get("size"): size_bits.append(r["size"])
            if r.get("jersey_number") and r["jersey_number"] not in ("X", ""): size_bits.append(f"#{r['jersey_number']}")
            size_str = " ".join(size_bits) or "-"
            qty = int(r.get("quantity") or 0); total_qty += qty
            delta = f"{r.get('qty_before', 0)}→{r.get('qty_after', 0)}"
            c.drawString(col_varenr, y, (r.get("full_number") or "-")[:13])
            c.drawString(col_dato, y, dato)
            c.drawString(col_vare, y, (r.get("item_name") or "-")[:28])
            c.drawString(col_str, y, size_str[:10])
            c.drawRightString(col_qty_r, y, str(qty))
            c.drawRightString(col_delta_r, y, delta[:14])
            # Overlapping range: ingen plass for både delta og order – skip order hvis tomt.
            if r.get("order_ref"):
                c.setFont("Helvetica", 7)
                c.drawString(col_order, y, (r.get("order_ref") or "")[:16])
                c.setFont("Helvetica", 8)
            c.drawRightString(col_user_r, y, (r.get("user_name") or "-")[:14])
            y -= 4 * mm

        y -= 2 * mm; c.line(LEFT, y, RIGHT, y); y -= 5 * mm
        c.setFont("Helvetica-Bold", 10)
        c.drawString(LEFT, y, f"Totalt mottatt: {total_qty} enheter fordelt på {len(receipts)} linjer")

    c.save()
    buf.seek(0)
    fn_date = datetime.now(timezone.utc).strftime("%Y%m%d")
    from fastapi.responses import StreamingResponse
    return StreamingResponse(buf, media_type="application/pdf",
                              headers={"Content-Disposition": f'attachment; filename="varemottak_{fn_date}.pdf"'})


# ============== TEAM NEEDS (Innmeldt Behov) ==============
NEED_RULES = {
    "Håndball":    {"min": 1, "max": 50, "size_required": True,  "jersey_required": False},
    "Hjemmedrakt": {"min": 1, "max": 1,  "size_required": True,  "jersey_required": True},
    "Bortedrakt":  {"min": 1, "max": 1,  "size_required": True,  "jersey_required": True},
    "Tape":        {"min": 1, "max": 10, "size_required": False, "jersey_required": False},
    "Isposer":     {"min": 1, "max": 10, "size_required": False, "jersey_required": False},
    "Andre klær":  {"min": 1, "max": 1,  "size_required": True,  "jersey_required": False},
}
NEED_HEADERS = ["Varetype", "Vare/variant", "Størrelse", "Draktnummer", "Antall", "Kommentar"]


def _current_seasons():
    y = datetime.now(timezone.utc).year
    return [f"{str(y)[-2:]}/{str(y+1)[-2:]}", f"{str(y+1)[-2:]}/{str(y+2)[-2:]}"]


@api.get("/needs/seasons")
async def need_seasons():
    return {"seasons": _current_seasons()}


async def _build_needs_template_bytes(tid: str, team_id: str) -> bytes:
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import PatternFill, Font
    team = await db.teams.find_one({"id": team_id, "tenant_id": tid}, {"_id": 0})
    if not team:
        raise HTTPException(404, "Lag finnes ikke")
    teams = await db.teams.find({"tenant_id": tid}, {"_id": 0}).to_list(500)
    items = await db.items.find({"tenant_id": tid}, {"_id": 0}).to_list(5000)
    seasons = _current_seasons()
    wb = Workbook()
    ws = wb.active; ws.title = "Behov"
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="7F1D1D", end_color="7F1D1D", fill_type="solid")
    ws["A1"] = "UTSTYRSROMMET – Innmeldt Behov"
    ws["A1"].font = Font(bold=True, size=14, color="7F1D1D")
    ws["A2"] = "Fyll ut lag, årsklasse og sesong, deretter én linje per enhet. Klær registreres som én enhet per linje."
    ws["A2"].font = Font(italic=True, size=10, color="555555")
    ws["A4"] = "Lagsnavn:"; ws["B4"] = team.get("team_name") or ""
    ws["A5"] = "Årsklasse:"; ws["B5"] = team.get("age_class") or ""
    ws["A6"] = "Sesong:"; ws["B6"] = seasons[0]
    ws["A7"] = "Lag-ID (ikke endre):"; ws["B7"] = team_id
    for r in range(4, 8):
        ws[f"A{r}"].font = Font(bold=True)
    for ci, h in enumerate(NEED_HEADERS, start=1):
        c = ws.cell(row=9, column=ci, value=h)
        c.font = hdr_font; c.fill = hdr_fill
    ref = wb.create_sheet("Referanse")
    team_names = sorted({t.get("team_name") for t in teams if t.get("team_name")})
    age_classes = sorted({t.get("age_class") for t in teams if t.get("age_class")})
    item_names = sorted({i.get("name") for i in items if i.get("name")})
    all_sizes = set()
    for it in items:
        for s in it.get("sizes") or []:
            if s.get("size"): all_sizes.add(s.get("size"))
    sizes = sorted(all_sizes)
    varetypes = list(NEED_RULES.keys())
    def _col(col_idx, header, values):
        ref.cell(row=1, column=col_idx, value=header)
        for i, v in enumerate(values, start=2):
            ref.cell(row=i, column=col_idx, value=v)
    _col(1, "Lag", team_names)
    _col(2, "Årsklasse", age_classes)
    _col(3, "Sesonger", seasons)
    _col(4, "Varetype", varetypes)
    _col(5, "Vare", item_names)
    _col(6, "Størrelse", sizes)
    dv_team = DataValidation(type="list", formula1=f"=Referanse!$A$2:$A${1+max(len(team_names),1)}", allow_blank=True)
    dv_age = DataValidation(type="list", formula1=f"=Referanse!$B$2:$B${1+max(len(age_classes),1)}", allow_blank=True)
    dv_season = DataValidation(type="list", formula1=f"=Referanse!$C$2:$C${1+len(seasons)}", allow_blank=True)
    dv_type = DataValidation(type="list", formula1=f"=Referanse!$D$2:$D${1+len(varetypes)}", allow_blank=False, errorTitle="Ugyldig varetype", error="Velg fra listen")
    dv_item = DataValidation(type="list", formula1=f"=Referanse!$E$2:$E${1+max(len(item_names),1)}", allow_blank=True)
    dv_size = DataValidation(type="list", formula1=f"=Referanse!$F$2:$F${1+max(len(sizes),1)}", allow_blank=True)
    dv_qty = DataValidation(type="whole", operator="between", formula1=1, formula2=50, allow_blank=True, errorTitle="Ugyldig antall", error="Antall må være heltall 1-50")
    for dv in (dv_team, dv_age, dv_season, dv_type, dv_item, dv_size, dv_qty):
        ws.add_data_validation(dv)
    dv_team.add("B4"); dv_age.add("B5"); dv_season.add("B6")
    rows_count = 205
    dv_type.add(f"A10:A{10+rows_count}")
    dv_item.add(f"B10:B{10+rows_count}")
    dv_size.add(f"C10:C{10+rows_count}")
    dv_qty.add(f"E10:E{10+rows_count}")
    widths = [16, 32, 14, 12, 9, 36]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64+i)].width = w
    ref.sheet_state = "hidden"
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


@api.get("/teams/{team_id}/needs/template")
async def team_needs_template(team_id: str, request: Request, user=Depends(require_perm("teams", "read"))):
    tid = await require_tenant(request, user)
    xlsx = await _build_needs_template_bytes(tid, team_id)
    team = await db.teams.find_one({"id": team_id, "tenant_id": tid}, {"_id": 0}) or {}
    seasons = _current_seasons()
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", team.get("team_name") or "lag")
    return Response(content=xlsx,
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="behov_mal_{safe}_{seasons[0].replace("/","-")}.xlsx"'})


class NeedConfirmIn(BaseModel):
    season: str
    lines: List[dict]
    strategy: str = "replace"
    filename: Optional[str] = None


async def _compute_need_coverage(tid, needs):
    items = await db.items.find({"tenant_id": tid}, {"_id": 0}).to_list(5000)
    item_by_id = {i["id"]: i for i in items}
    def vk(iid, size, jersey): return f"{iid}|{size or ''}|{jersey or ''}"
    avail = {}
    for i in items:
        if i.get("is_clothing") or i.get("is_ball"):
            for s in i.get("sizes") or []:
                jk = s.get("jersey_number") if i.get("is_clothing") else ""
                avail[vk(i["id"], s.get("size"), jk)] = int(s.get("quantity") or 0)
        else:
            avail[vk(i["id"], None, None)] = int(i.get("total_stock") or 0)
    computed = []
    for need in needs:
        for ln in need.get("lines", []):
            item_id = ln.get("item_id")
            size = ln.get("size") or None
            jersey = ln.get("jersey_number") or ""
            it = item_by_id.get(item_id, {})
            key = vk(item_id, size, jersey if it.get("is_clothing") else "")
            avail_now = avail.get(key, 0)
            needed = int(ln.get("quantity") or 0)
            can_cover = min(needed, avail_now)
            avail[key] = max(0, avail_now - can_cover)
            must_order = max(0, needed - can_cover)
            pprice = it.get("purchase_price")
            missing_price = pprice in (None, "", 0)
            cost = (must_order * pprice) if not missing_price else None
            alt = None
            if must_order > 0 and it:
                for cand in items:
                    if cand.get("id") == item_id: continue
                    if bool(cand.get("is_clothing")) != bool(it.get("is_clothing")): continue
                    if bool(cand.get("is_ball")) != bool(it.get("is_ball")): continue
                    if (cand.get("category") or "") != (it.get("category") or ""): continue
                    if cand.get("is_clothing") or cand.get("is_ball"):
                        for cs in cand.get("sizes") or []:
                            if cs.get("size") == size and int(cs.get("quantity") or 0) > 0:
                                alt = {"item_id": cand["id"], "name": cand.get("name"), "size": cs.get("size"),
                                       "quantity_available": int(cs.get("quantity") or 0),
                                       "reason": "Samme kategori og størrelse"}
                                break
                    else:
                        if int(cand.get("total_stock") or 0) > 0:
                            alt = {"item_id": cand["id"], "name": cand.get("name"), "size": None,
                                   "quantity_available": int(cand.get("total_stock") or 0),
                                   "reason": "Samme kategori"}
                    if alt: break
            computed.append({
                "need_id": need.get("id"), "line_id": ln.get("id"),
                "item_type": ln.get("item_type"), "item_id": item_id,
                "item_name": it.get("name") or ln.get("item_name") or "?",
                "size": size, "jersey_number": jersey, "quantity": needed,
                "available_now": avail_now, "can_cover": can_cover, "must_order": must_order,
                "unit_purchase_price": pprice, "line_cost": cost,
                "missing_price": missing_price, "alternative": alt,
                "comment": ln.get("comment"),
            })
    return computed


@api.get("/teams/{team_id}/needs")
async def get_team_needs(team_id: str, request: Request, season: Optional[str] = None, user=Depends(require_perm("teams", "read"))):
    tid = await require_tenant(request, user)
    q = {"tenant_id": tid, "team_id": team_id}
    if season: q["season"] = season
    needs = await db.team_needs.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)
    computed = await _compute_need_coverage(tid, needs)
    return {"seasons": _current_seasons(), "needs": needs, "lines_computed": computed,
            "computed_at": now_iso()}


@api.post("/teams/{team_id}/needs/import")
async def import_team_needs(team_id: str, request: Request, file: UploadFile = File(...), user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    team = await db.teams.find_one({"id": team_id, "tenant_id": tid}, {"_id": 0})
    if not team:
        raise HTTPException(404, "Lag finnes ikke")
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(400, "Kunne ikke lese Excel-fil")
    ws = wb.active
    file_season = str(ws["B6"].value or "").strip()
    file_team_id = str(ws["B7"].value or "").strip()
    file_team_name = str(ws["B4"].value or "").strip()
    file_age = str(ws["B5"].value or "").strip()
    if file_team_id and file_team_id != team_id:
        raise HTTPException(400, f"Malen tilhører et annet lag ({file_team_name!r}).")
    valid_seasons = _current_seasons()
    if file_season not in valid_seasons:
        raise HTTPException(400, f"Ugyldig sesong: {file_season!r}. Gyldige: {', '.join(valid_seasons)}")
    items = await db.items.find({"tenant_id": tid}, {"_id": 0}).to_list(5000)
    items_by_name = {i["name"].lower(): i for i in items if i.get("name")}
    lines_preview = []
    errors = []
    for ridx, row in enumerate(ws.iter_rows(min_row=10, values_only=True), start=10):
        if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in (row[:6] if row else [])):
            continue
        vtype = (str(row[0]) if row[0] else "").strip()
        iname = (str(row[1]) if row[1] else "").strip()
        size = (str(row[2]) if row[2] else "").strip()
        jersey = (str(row[3]) if row[3] else "").strip()
        qty_raw = row[4] if len(row) > 4 else None
        comment = (str(row[5]) if len(row) > 5 and row[5] else "").strip()
        line_err = []
        if vtype not in NEED_RULES:
            line_err.append(f"Ukjent varetype {vtype!r}")
        rule = NEED_RULES.get(vtype, {})
        try:
            qty = int(qty_raw) if qty_raw not in (None, "") else 0
        except Exception:
            qty = 0; line_err.append("Antall må være heltall")
        if rule and (qty < rule.get("min", 1) or qty > rule.get("max", 50)):
            line_err.append(f"Antall {qty} ugyldig for {vtype} (må være {rule.get('min')}-{rule.get('max')})")
        if rule.get("size_required") and not size:
            line_err.append("Størrelse er påkrevd")
        if rule.get("jersey_required") and not jersey:
            line_err.append("Draktnummer er påkrevd")
        it = items_by_name.get(iname.lower()) if iname else None
        if not iname:
            line_err.append("Vare/variant mangler")
        elif not it:
            line_err.append(f"Fant ikke vare {iname!r} i katalogen")
        lines_preview.append({
            "row": ridx, "item_type": vtype, "item_name": iname,
            "item_id": it["id"] if it else None,
            "size": size, "jersey_number": jersey,
            "quantity": qty, "comment": comment, "errors": line_err,
        })
        if line_err:
            errors.extend([f"Rad {ridx}: {e}" for e in line_err])
    existing = await db.team_needs.find({"tenant_id": tid, "team_id": team_id, "season": file_season}, {"_id": 0}).to_list(100)
    return {
        "team_id": team_id, "team_name": team.get("team_name"),
        "age_class": file_age, "season": file_season,
        "filename": file.filename, "lines": lines_preview,
        "errors": errors, "has_existing": bool(existing),
    }


@api.post("/teams/{team_id}/needs/confirm")
async def confirm_team_needs(team_id: str, data: NeedConfirmIn, request: Request, user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    if data.season not in _current_seasons():
        raise HTTPException(400, "Ugyldig sesong")
    clean = []
    for ln in data.lines:
        if ln.get("errors"): continue
        if not ln.get("item_id"): continue
        clean.append({
            "id": new_id(),
            "item_type": ln.get("item_type"),
            "item_id": ln.get("item_id"),
            "item_name": ln.get("item_name"),
            "size": ln.get("size") or None,
            "jersey_number": ln.get("jersey_number") or None,
            "quantity": int(ln.get("quantity") or 0),
            "comment": ln.get("comment") or "",
            "original_item_id": ln.get("item_id"),
        })
    if not clean:
        raise HTTPException(400, "Ingen gyldige linjer å importere")
    if data.strategy == "replace":
        await db.team_needs.delete_many({"tenant_id": tid, "team_id": team_id, "season": data.season})
    doc = {
        "id": new_id(), "tenant_id": tid, "team_id": team_id, "season": data.season,
        "filename": data.filename or "", "lines": clean,
        "imported_by": user.get("name") or user.get("email"),
        "created_at": now_iso(),
    }
    await db.team_needs.insert_one(doc)
    return {"ok": True, "need_id": doc["id"]}


@api.delete("/teams/{team_id}/needs/{need_id}")
async def delete_team_need(team_id: str, need_id: str, request: Request, user=Depends(require_perm("teams", "admin"))):
    tid = await require_tenant(request, user)
    r = await db.team_needs.delete_one({"id": need_id, "tenant_id": tid, "team_id": team_id})
    if not r.deleted_count:
        raise HTTPException(404, "Fant ikke")
    return {"ok": True}


# ============== LEVERANDØRER ==============
class SupplierIn(BaseModel):
    name: str
    address: str = ""
    postal_code: str = ""
    city: str = ""
    country: str = "Norge"
    email: str = ""
    phone: str = ""
    contact_person: str = ""
    org_number: str = ""
    notes: str = ""
    active: bool = True


async def _next_supplier_number(tid: str) -> str:
    key = f"supplier_counter_{tid}"
    doc = await db.counters.find_one_and_update(
        {"_id": key}, {"$inc": {"seq": 1}}, upsert=True, return_document=True
    )
    seq = doc.get("seq", 1) if doc else 1
    return f"{seq:05d}"


@api.get("/suppliers")
async def list_suppliers(request: Request, user=Depends(get_current_user)):
    tid = await require_tenant(request, user)
    not_deleted = {"$or": [{"deleted_at": None}, {"deleted_at": {"$exists": False}}]}
    # Backfill any suppliers missing supplier_number (men ikke slettede)
    missing = await db.suppliers.find({"tenant_id": tid, "supplier_number": {"$in": [None, ""]}, **not_deleted}, {"_id": 0}).to_list(500)
    for m in missing:
        num = await _next_supplier_number(tid)
        await db.suppliers.update_one({"id": m["id"], "tenant_id": tid}, {"$set": {"supplier_number": num}})
    return await db.suppliers.find({"tenant_id": tid, **not_deleted}, {"_id": 0}).sort("name", 1).to_list(500)


@api.get("/suppliers/trash")
async def list_deleted_suppliers(request: Request, user=Depends(require_admin)):
    """Soft-slettede leverandører fra siste 30 dager. Kun admin."""
    tid = await require_tenant(request, user)
    cutoff = (datetime.now(timezone.utc) - timedelta(days=30)).isoformat()
    docs = await db.suppliers.find({
        "tenant_id": tid,
        "deleted_at": {"$ne": None, "$gte": cutoff},
    }, {"_id": 0}).sort("deleted_at", -1).to_list(200)
    return docs


@api.post("/suppliers/{sid}/restore")
async def restore_supplier(sid: str, request: Request, user=Depends(require_admin)):
    """Gjenopprett en soft-slettet leverandør. Kun admin."""
    tid = await require_tenant(request, user)
    doc = await db.suppliers.find_one({"id": sid, "tenant_id": tid})
    if not doc:
        raise HTTPException(404, "Fant ikke leverandør")
    if not doc.get("deleted_at"):
        raise HTTPException(400, "Leverandøren er ikke slettet")
    await db.suppliers.update_one({"id": sid, "tenant_id": tid}, {
        "$unset": {"deleted_at": "", "deleted_by_user_id": "", "deleted_by_name": ""},
        "$set": {"active": True, "restored_at": now_iso(), "restored_by_name": user.get("name") or user.get("email")},
    })
    return await db.suppliers.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})


@api.post("/suppliers")
async def create_supplier(data: SupplierIn, request: Request, user=Depends(require_admin)):
    tid = await require_tenant(request, user)
    if not data.name.strip():
        raise HTTPException(400, "Leverandørnavn er obligatorisk")
    num = await _next_supplier_number(tid)
    doc = {**data.dict(), "id": new_id(), "tenant_id": tid,
           "supplier_number": num, "created_at": now_iso()}
    await db.suppliers.insert_one(doc); doc.pop("_id", None); return doc


@api.put("/suppliers/{sid}")
async def update_supplier(sid: str, data: SupplierIn, request: Request, user=Depends(require_admin)):
    tid = await require_tenant(request, user)
    await db.suppliers.update_one({"id": sid, "tenant_id": tid}, {"$set": data.dict()})
    return await db.suppliers.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})


@api.post("/suppliers/{sid}/toggle-active")
async def toggle_supplier_active(sid: str, request: Request, user=Depends(require_admin)):
    """Veksler active-status for leverandør. Deaktiverte leverandører skjules fra nye bestillinger,
    men beholder all historikk og vises fortsatt i leverandørlisten."""
    tid = await require_tenant(request, user)
    doc = await db.suppliers.find_one({"id": sid, "tenant_id": tid})
    if not doc:
        raise HTTPException(404, "Fant ikke leverandør")
    if doc.get("deleted_at"):
        raise HTTPException(400, "Kan ikke endre status på en slettet leverandør – gjenopprett først")
    new_active = not bool(doc.get("active", True))
    await db.suppliers.update_one({"id": sid, "tenant_id": tid}, {"$set": {
        "active": new_active,
        "active_changed_at": now_iso(),
        "active_changed_by": user.get("name") or user.get("email"),
    }})
    return {"ok": True, "active": new_active}


@api.delete("/suppliers/{sid}")
async def delete_supplier(sid: str, request: Request, user=Depends(require_admin)):
    """Soft-delete leverandør. Admin kan angre innen 30 dager via POST /suppliers/{sid}/restore."""
    tid = await require_tenant(request, user)
    doc = await db.suppliers.find_one({"id": sid, "tenant_id": tid})
    if not doc:
        raise HTTPException(404, "Fant ikke leverandør")
    orders = await db.purchase_orders.count_documents({"tenant_id": tid, "supplier_id": sid})
    await db.suppliers.update_one({"id": sid, "tenant_id": tid}, {"$set": {
        "deleted_at": now_iso(),
        "deleted_by_user_id": user.get("id"),
        "deleted_by_name": user.get("name") or user.get("email"),
        "active": False,
    }})
    return {"ok": True, "soft_deleted": True, "had_orders": orders > 0}


@api.get("/suppliers/{sid}/orders")
async def supplier_orders(sid: str, request: Request, user=Depends(get_current_user)):
    tid = await require_tenant(request, user)
    return await db.purchase_orders.find({"tenant_id": tid, "supplier_id": sid}, {"_id": 0}).sort("created_at", -1).to_list(500)


# ============== UNIQUE PO NUMBER ==============
async def _next_po_number(tid: str, prefix: str = "") -> str:
    """5-digit sequence per tenant per year, with -YY suffix. Thread-safe."""
    yr = datetime.now(timezone.utc).strftime("%y")
    key = f"po_{tid}_{yr}"
    doc = await db.counters.find_one_and_update(
        {"_id": key}, {"$inc": {"seq": 1}}, upsert=True, return_document=True
    )
    seq = doc.get("seq", 1) if doc else 1
    return f"{prefix}{seq:05d}-{yr}"


# ============== BEHOVSSKJEMA ==============
class BehovsskjemaIn(BaseModel):
    team_id: str
    season: str


@api.post("/behovsskjema")
async def create_behovsskjema(data: BehovsskjemaIn, request: Request, user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    team = await db.teams.find_one({"id": data.team_id, "tenant_id": tid}, {"_id": 0})
    if not team:
        raise HTTPException(404, "Lag finnes ikke")
    if data.season not in _current_seasons():
        raise HTTPException(400, f"Ugyldig sesong. Gyldige: {', '.join(_current_seasons())}")
    po_num = await _next_po_number(tid)
    doc = {
        "id": new_id(), "tenant_id": tid, "po_number": po_num,
        "team_id": data.team_id, "team_name": team.get("team_name"),
        "age_class": team.get("age_class"), "season": data.season,
        "status": "ikke_sendt", "sent_at": None, "sent_to": None,
        "imported_at": None, "last_import_file": None,
        "notes": "", "created_at": now_iso(),
        "created_by": user.get("name") or user.get("email"),
    }
    await db.behovsskjema.insert_one(doc); doc.pop("_id", None)
    return doc


@api.get("/behovsskjema")
async def list_behovsskjema(request: Request, user=Depends(require_perm("teams", "read")), team_id: Optional[str] = None):
    tid = await require_tenant(request, user)
    q = {"tenant_id": tid}
    if team_id: q["team_id"] = team_id
    return await db.behovsskjema.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)


@api.put("/behovsskjema/{sid}")
async def update_behovsskjema(sid: str, data: dict, request: Request, user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    allowed = {k: data[k] for k in ("notes", "status", "sent_at", "sent_to",
                                     "recipient_name", "recipient_email", "message") if k in data}
    await db.behovsskjema.update_one({"id": sid, "tenant_id": tid}, {"$set": allowed})
    return await db.behovsskjema.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})


@api.post("/behovsskjema/{sid}/send")
async def send_behovsskjema(sid: str, data: dict, request: Request, user=Depends(require_perm("teams", "write"))):
    """Send behovsskjema (Excel template) via email to lagleder."""
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})
    if not bs:
        raise HTTPException(404, "Fant ikke")
    to_email = (data.get("recipient_email") or bs.get("recipient_email") or "").strip()
    if not to_email or not re.match(r"^[^@]+@[^@]+\.[^@]+$", to_email):
        raise HTTPException(400, "Ugyldig e-postadresse")
    if not EMAIL_KEY:
        raise HTTPException(503, "E-posttjeneste er ikke konfigurert")
    # Generate Excel template directly as bytes
    try:
        xlsx_bytes = await _build_needs_template_bytes(tid, bs["team_id"])
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Template generation failed")
        raise HTTPException(500, f"Kunne ikke generere mal: {e}")
    import base64
    att_b64 = base64.b64encode(xlsx_bytes).decode()
    recipient_name = data.get("recipient_name") or bs.get("recipient_name") or ""
    custom_msg = data.get("message") or bs.get("message") or ""
    club = await db.customers.find_one({"id": tid}, {"_id": 0}) or {}
    html = (
        f'<p>Hei {recipient_name or ""},</p>'
        f'<p>Her er behovsskjemaet for <strong>{bs.get("team_name")}</strong> '
        f'(årsklasse {bs.get("age_class")}, sesong {bs.get("season")}).</p>'
        f'<p><strong>Innkjøpsnummer:</strong> {bs.get("po_number")}</p>'
        + (f'<p>{custom_msg}</p>' if custom_msg else '')
        + f'<p>Vennligst fyll ut Excel-malen og returner den til oss.</p>'
        f'<p>Mvh<br/>{club.get("club_name") or "Klubben"}</p>'
    )
    filename = f"behovsskjema_{bs.get('po_number')}_{(bs.get('team_name') or '').replace(' ','_')}.xlsx"
    attachments = [{
        "filename": filename,
        "content": att_b64,
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }]
    cc_list = [user.get("email")] if user.get("email") else None
    res = await send_tenant_email(
        tid, to_email,
        f"Behovsskjema {bs.get('po_number')} – {bs.get('team_name')}",
        html, attachments=attachments, cc=cc_list,
        sender_user=user,
    )
    if not res.get("ok"):
        raise HTTPException(502, f"E-post feilet: {res.get('error')}")
    provider_id = res.get("provider_id")
    await db.behovsskjema.update_one({"id": sid, "tenant_id": tid}, {"$set": {
        "status": "sendt", "sent_at": now_iso(), "sent_to": to_email,
        "recipient_name": recipient_name, "recipient_email": to_email,
        "last_email_id": provider_id,
    }})
    return {"ok": True, "sent_to": to_email, "provider_id": provider_id, "method": res.get("method")}


@api.delete("/behovsskjema/{sid}")
async def delete_behovsskjema(sid: str, request: Request, user=Depends(require_perm("teams", "admin"))):
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": sid, "tenant_id": tid})
    if not bs:
        raise HTTPException(404, "Fant ikke")
    orders = await db.purchase_orders.count_documents({"tenant_id": tid, "behovsskjema_id": sid})
    if orders:
        await db.behovsskjema.update_one({"id": sid, "tenant_id": tid}, {"$set": {"status": "arkivert"}})
        return {"ok": True, "archived": True}
    await db.behovsskjema.delete_one({"id": sid, "tenant_id": tid})
    return {"ok": True}


@api.post("/behovsskjema/{sid}/import")
async def import_behovsskjema(sid: str, file: UploadFile = File(...), request: Request = None,
                                user=Depends(require_perm("teams", "write"))):
    """Importer utfylt Excel-behovsskjema tilbake. Parser linjer og lagrer som draft_lines
    som kan redigeres i app før bekreftelse. Status settes til 'importert'."""
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})
    if not bs:
        raise HTTPException(404, "Fant ikke")
    if bs.get("status") == "bekreftet":
        raise HTTPException(409, "Behovsskjema er allerede bekreftet – kan ikke importeres på nytt. Slett bekreftelsen først.")
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(400, "Kunne ikke lese Excel-fil")
    ws = wb.active
    file_team_id = str(ws["B7"].value or "").strip()
    file_season = str(ws["B6"].value or "").strip()
    if file_team_id and file_team_id != bs["team_id"]:
        raise HTTPException(400, f"Filen tilhører et annet lag enn behovsskjemaet. Last ned riktig mal på nytt.")
    items = await db.items.find({"tenant_id": tid}, {"_id": 0}).to_list(5000)
    items_by_name = {i["name"].lower(): i for i in items if i.get("name")}
    draft_lines = []; errors = []
    for ridx, row in enumerate(ws.iter_rows(min_row=10, values_only=True), start=10):
        if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in (row[:6] if row else [])):
            continue
        vtype = (str(row[0]) if row[0] else "").strip()
        iname = (str(row[1]) if row[1] else "").strip()
        size = (str(row[2]) if row[2] else "").strip()
        jersey = (str(row[3]) if row[3] else "").strip()
        qty_raw = row[4] if len(row) > 4 else None
        comment = (str(row[5]) if len(row) > 5 and row[5] else "").strip()
        line_err = []
        if vtype and vtype not in NEED_RULES:
            line_err.append(f"Ukjent varetype «{vtype}»")
        rule = NEED_RULES.get(vtype, {})
        try:
            qty = int(qty_raw) if qty_raw not in (None, "") else 0
        except Exception:
            qty = 0; line_err.append("Antall må være heltall")
        if rule and (qty < rule.get("min", 1) or qty > rule.get("max", 50)):
            line_err.append(f"Antall {qty} ugyldig for {vtype} (må være {rule.get('min')}-{rule.get('max')})")
        if rule.get("size_required") and not size:
            line_err.append("Størrelse er påkrevd")
        if rule.get("jersey_required") and not jersey:
            line_err.append("Draktnummer er påkrevd")
        it = items_by_name.get(iname.lower()) if iname else None
        if not iname:
            line_err.append("Vare mangler")
        elif not it:
            line_err.append(f"Fant ikke vare «{iname}» i katalogen")
        draft_lines.append({
            "line_id": new_id(),
            "row": ridx, "item_type": vtype, "item_name": iname,
            "item_id": it["id"] if it else None,
            "size": size or None, "jersey_number": jersey or None,
            "quantity": qty, "comment": comment, "errors": line_err,
        })
        if line_err:
            errors.extend([f"Rad {ridx}: {e}" for e in line_err])
    await db.behovsskjema.update_one({"id": sid, "tenant_id": tid}, {"$set": {
        "status": "importert", "imported_at": now_iso(),
        "last_import_file": file.filename, "draft_lines": draft_lines,
        "import_season": file_season or bs.get("season"),
    }})
    if user.get("role") != "platform_owner":
        await _log_activity(tid, user, "import_behovsskjema",
                             f"Importerte utfylt behovsskjema {bs.get('po_number')} ({len(draft_lines)} linjer, {len(errors)} feil)",
                             {"sid": sid, "lines": len(draft_lines)})
    return {"ok": True, "lines": draft_lines, "errors": errors, "total": len(draft_lines)}


@api.put("/behovsskjema/{sid}/lines")
async def update_behovsskjema_lines(sid: str, data: dict, request: Request,
                                     user=Depends(require_perm("teams", "write"))):
    """Erstatt draft_lines med redigert liste. Lar bruker endre antall, str, draktnr, kommentar, vare."""
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})
    if not bs:
        raise HTTPException(404, "Fant ikke")
    if bs.get("status") == "bekreftet":
        raise HTTPException(409, "Behovsskjema er bekreftet – kan ikke redigeres. Slett bekreftelsen først.")
    raw_lines = data.get("lines") or []
    items = await db.items.find({"tenant_id": tid}, {"_id": 0}).to_list(5000)
    items_by_id = {i["id"]: i for i in items}
    items_by_name = {i["name"].lower(): i for i in items if i.get("name")}
    clean = []
    for ln in raw_lines:
        it = None
        if ln.get("item_id"):
            it = items_by_id.get(ln["item_id"])
        if not it and ln.get("item_name"):
            it = items_by_name.get(ln["item_name"].lower())
        line_err = []
        qty = int(ln.get("quantity") or 0)
        vtype = (ln.get("item_type") or "").strip()
        rule = NEED_RULES.get(vtype, {})
        if vtype and vtype not in NEED_RULES:
            line_err.append(f"Ukjent varetype «{vtype}»")
        if rule and (qty < rule.get("min", 1) or qty > rule.get("max", 50)):
            line_err.append(f"Antall {qty} ugyldig for {vtype}")
        if rule.get("size_required") and not ln.get("size"):
            line_err.append("Størrelse er påkrevd")
        if rule.get("jersey_required") and not ln.get("jersey_number"):
            line_err.append("Draktnummer er påkrevd")
        if not it:
            line_err.append("Vare finnes ikke i katalogen")
        clean.append({
            "line_id": ln.get("line_id") or new_id(),
            "item_type": vtype, "item_name": it.get("name") if it else (ln.get("item_name") or ""),
            "item_id": it["id"] if it else None,
            "size": (ln.get("size") or "").strip() or None,
            "jersey_number": (str(ln.get("jersey_number") or "")).strip() or None,
            "quantity": qty, "comment": (ln.get("comment") or "").strip(),
            "errors": line_err,
        })
    await db.behovsskjema.update_one({"id": sid, "tenant_id": tid}, {"$set": {"draft_lines": clean}})
    return {"ok": True, "lines": clean, "total": len(clean), "errors_count": sum(1 for l in clean if l["errors"])}


@api.post("/behovsskjema/{sid}/confirm")
async def confirm_behovsskjema(sid: str, request: Request, user=Depends(require_perm("teams", "write"))):
    """Bekreft behovsskjema som fullført behov for laget. Oppretter team_needs-rad og
    låser behovsskjemaet med status='bekreftet'. Krever at alle linjer er feilfrie."""
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})
    if not bs:
        raise HTTPException(404, "Fant ikke")
    if bs.get("status") == "bekreftet":
        raise HTTPException(409, "Behovsskjema er allerede bekreftet")
    draft_lines = bs.get("draft_lines") or []
    if not draft_lines:
        raise HTTPException(400, "Ingen linjer å bekrefte – importer utfylt skjema først")
    err_count = sum(1 for l in draft_lines if l.get("errors"))
    if err_count:
        raise HTTPException(400, f"{err_count} linjer har valideringsfeil – fiks dem før du bekrefter")
    season = bs.get("import_season") or bs.get("season")
    # Fjern eventuelle tidligere bekreftede behov for samme lag+sesong
    await db.team_needs.delete_many({"tenant_id": tid, "team_id": bs["team_id"], "season": season,
                                      "behovsskjema_id": sid})
    confirmed_lines = []
    for ln in draft_lines:
        confirmed_lines.append({
            "id": new_id(),
            "item_type": ln.get("item_type"),
            "item_id": ln.get("item_id"),
            "item_name": ln.get("item_name"),
            "size": ln.get("size"),
            "jersey_number": ln.get("jersey_number"),
            "quantity": int(ln.get("quantity") or 0),
            "comment": ln.get("comment") or "",
            "original_item_id": ln.get("item_id"),
        })
    doc = {
        "id": new_id(), "tenant_id": tid, "team_id": bs["team_id"], "season": season,
        "behovsskjema_id": sid, "behovsskjema_po": bs.get("po_number"),
        "filename": bs.get("last_import_file") or "", "lines": confirmed_lines,
        "imported_by": user.get("name") or user.get("email"), "created_at": now_iso(),
    }
    await db.team_needs.insert_one(doc)
    await db.behovsskjema.update_one({"id": sid, "tenant_id": tid}, {"$set": {
        "status": "bekreftet", "confirmed_at": now_iso(),
        "confirmed_by": user.get("name") or user.get("email"),
        "team_need_id": doc["id"],
    }})
    if user.get("role") != "platform_owner":
        await _log_activity(tid, user, "confirm_behovsskjema",
                             f"Bekreftet behovsskjema {bs.get('po_number')} som behov for {bs.get('team_name')} ({len(confirmed_lines)} linjer)",
                             {"sid": sid, "team_need_id": doc["id"]})
    return {"ok": True, "team_need_id": doc["id"], "lines_count": len(confirmed_lines)}


@api.post("/behovsskjema/{sid}/unlock")
async def unlock_behovsskjema(sid: str, request: Request, user=Depends(require_perm("teams", "admin"))):
    """Fjern bekreftelsen slik at linjene kan redigeres igjen. Fjerner tilhørende team_needs-rad."""
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})
    if not bs:
        raise HTTPException(404, "Fant ikke")
    if bs.get("status") != "bekreftet":
        raise HTTPException(400, "Behovsskjema er ikke bekreftet")
    if bs.get("team_need_id"):
        await db.team_needs.delete_one({"id": bs["team_need_id"], "tenant_id": tid})
    await db.behovsskjema.update_one({"id": sid, "tenant_id": tid}, {"$set": {
        "status": "importert", "confirmed_at": None, "confirmed_by": None, "team_need_id": None,
    }})
    return {"ok": True}


# ============== PURCHASE ORDERS (LEVERANDØRBESTILLINGER) ==============
class OrderLineIn(BaseModel):
    item_id: Optional[str] = None
    item_name: str
    size: Optional[str] = None
    jersey_number: Optional[str] = None
    quantity: int
    unit_price: float = 0.0


class PurchaseOrderIn(BaseModel):
    supplier_id: str
    behovsskjema_id: Optional[str] = None  # None for lagerbestilling
    team_id: Optional[str] = None  # None for lagerbestilling
    season: Optional[str] = None
    is_stock_order: bool = False
    lines: List[OrderLineIn]
    notes: str = ""
    delivery_date: Optional[str] = None


@api.post("/orders")
async def create_order(data: PurchaseOrderIn, request: Request, user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    supplier = await db.suppliers.find_one({"id": data.supplier_id, "tenant_id": tid}, {"_id": 0})
    if not supplier:
        raise HTTPException(404, "Leverandør finnes ikke")
    if not data.lines:
        raise HTTPException(400, "Minst én varelinje kreves")
    # Lagerbestilling: new LB-prefix number. Team order: reuse behovsskjema's po_number.
    po_number = None
    behov = None
    if data.is_stock_order:
        po_number = await _next_po_number(tid, prefix="LB")
    else:
        if not data.behovsskjema_id:
            raise HTTPException(400, "behovsskjema_id kreves for lagsbestilling")
        behov = await db.behovsskjema.find_one({"id": data.behovsskjema_id, "tenant_id": tid}, {"_id": 0})
        if not behov:
            raise HTTPException(404, "Behovsskjema finnes ikke")
        po_number = behov["po_number"]
    total = sum((ln.quantity or 0) * (ln.unit_price or 0) for ln in data.lines)
    order_lines = []
    for ln in data.lines:
        order_lines.append({
            "id": new_id(), **ln.dict(),
            "received_quantity": 0, "line_status": "ulevert",
            "receipts": [], "reserved_from_po": None, "reserved_for_po": None,
        })
    doc = {
        "id": new_id(), "tenant_id": tid, "po_number": po_number,
        "supplier_id": data.supplier_id, "supplier_name": supplier.get("name"),
        "behovsskjema_id": data.behovsskjema_id,
        "team_id": data.team_id or (behov["team_id"] if behov else None),
        "team_name": behov["team_name"] if behov else None,
        "age_class": behov["age_class"] if behov else None,
        "season": data.season or (behov["season"] if behov else None),
        "is_stock_order": data.is_stock_order,
        "lines": order_lines, "notes": data.notes,
        "delivery_date": data.delivery_date,
        "status": "bestilt",  # bestilt | delvis_levert | levert | manko | restbestilling | arkivert
        "manual_status": None, "manual_status_reason": None,
        "status_history": [{"status": "bestilt", "at": now_iso(), "by": user.get("name"), "note": "Opprettet"}],
        "pdf_versions": [],
        "sent_at": None, "sent_to": None,
        "total_amount": total,
        "created_at": now_iso(), "created_by": user.get("name") or user.get("email"),
        "ready_for_pickup_notified": False,
    }
    await db.purchase_orders.insert_one(doc)
    if data.behovsskjema_id:
        await db.behovsskjema.update_one({"id": data.behovsskjema_id, "tenant_id": tid},
                                         {"$set": {"status": "bestilt"}})
    doc.pop("_id", None)
    return doc


@api.get("/orders")
async def list_orders(request: Request, user=Depends(get_current_user),
                      team_id: Optional[str] = None, supplier_id: Optional[str] = None,
                      is_stock_order: Optional[bool] = None):
    tid = await require_tenant(request, user)
    q = {"tenant_id": tid}
    if team_id: q["team_id"] = team_id
    if supplier_id: q["supplier_id"] = supplier_id
    if is_stock_order is not None: q["is_stock_order"] = is_stock_order
    return await db.purchase_orders.find(q, {"_id": 0}).sort("created_at", -1).to_list(500)


@api.get("/orders/by-po/{po_number}")
async def get_order_by_po(po_number: str, request: Request, user=Depends(get_current_user)):
    tid = await require_tenant(request, user)
    o = await db.purchase_orders.find_one({"tenant_id": tid, "po_number": po_number}, {"_id": 0})
    if not o:
        raise HTTPException(404, "Fant ingen bestilling med dette nummeret")
    return o


class OrderReceiveLine(BaseModel):
    line_id: str
    quantity: int
    comment: str = ""
    confirm_over: bool = False


class OrderReceiveIn(BaseModel):
    receipts: List[OrderReceiveLine]
    request_id: Optional[str] = None


@api.post("/orders/{oid}/receive")
async def receive_order(oid: str, data: OrderReceiveIn, request: Request, user=Depends(require_perm("inventory", "write"))):
    tid = await require_tenant(request, user)
    # idempotency
    if data.request_id:
        existing = await db.stock_receipts.find_one({"tenant_id": tid, "request_id": data.request_id}, {"_id": 0})
        if existing:
            return {"ok": True, "idempotent": True}
    order = await db.purchase_orders.find_one({"id": oid, "tenant_id": tid})
    if not order:
        raise HTTPException(404, "Bestilling finnes ikke")
    lines = order.get("lines") or []
    updates = []
    for rec in data.receipts:
        if rec.quantity <= 0: continue
        ln = next((l for l in lines if l["id"] == rec.line_id), None)
        if not ln:
            raise HTTPException(400, f"Fant ikke linje {rec.line_id}")
        remaining = (ln.get("quantity") or 0) - (ln.get("received_quantity") or 0)
        if rec.quantity > remaining and not rec.confirm_over:
            raise HTTPException(400, f"Mottatt {rec.quantity} er mer enn restantall {remaining}. Bekreft med confirm_over.")
        ln["received_quantity"] = (ln.get("received_quantity") or 0) + rec.quantity
        ln["receipts"] = (ln.get("receipts") or []) + [{
            "at": now_iso(), "by": user.get("name"), "quantity": rec.quantity, "comment": rec.comment,
        }]
        if ln["received_quantity"] >= (ln.get("quantity") or 0):
            ln["line_status"] = "levert"; ln["delivered_at"] = now_iso()
        else:
            ln["line_status"] = "delvis_levert"
        # Update inventory
        if ln.get("item_id"):
            item = await db.items.find_one({"id": ln["item_id"], "tenant_id": tid})
            if item:
                if item.get("is_clothing") or item.get("is_ball"):
                    sizes = item.get("sizes") or []
                    for s in sizes:
                        if s.get("size") == ln.get("size"):
                            s["quantity"] = int(s.get("quantity") or 0) + rec.quantity
                            break
                    new_total = sum(int(s.get("quantity") or 0) for s in sizes)
                    await db.items.update_one({"id": ln["item_id"], "tenant_id": tid},
                                              {"$set": {"sizes": sizes, "total_stock": new_total}})
                else:
                    await db.items.update_one({"id": ln["item_id"], "tenant_id": tid},
                                              {"$inc": {"total_stock": rec.quantity}})
        # Also log in stock_receipts for history
        await db.stock_receipts.insert_one({
            "id": new_id(), "tenant_id": tid, "request_id": data.request_id,
            "item_id": ln.get("item_id"), "item_name": ln.get("item_name"),
            "full_number": ln.get("full_number") or ln.get("item_number") or "",
            "size": ln.get("size"), "jersey_number": ln.get("jersey_number"),
            "quantity": rec.quantity, "qty_before": 0, "qty_after": 0,
            "order_ref": order.get("po_number"), "comment": rec.comment,
            "order_id": oid, "order_line_id": ln["id"],
            "user_id": user.get("id"), "user_name": user.get("name"),
            "created_at": now_iso(),
        })
        updates.append(ln["id"])
    # Order status
    all_lines = lines
    total_q = sum(l.get("quantity") or 0 for l in all_lines)
    recv_q = sum(l.get("received_quantity") or 0 for l in all_lines)
    new_status = "bestilt"
    if recv_q >= total_q and total_q > 0: new_status = "levert"
    elif recv_q > 0: new_status = "delvis_levert"
    status_entry = {"status": new_status, "at": now_iso(), "by": user.get("name"), "note": f"Mottak {len(updates)} linjer"}
    await db.purchase_orders.update_one({"id": oid, "tenant_id": tid},
                                        {"$set": {"lines": lines, "status": new_status},
                                         "$push": {"status_history": status_entry}})
    return {"ok": True, "status": new_status, "lines_updated": updates}


class ManualStatusIn(BaseModel):
    status: str
    reason: str


@api.post("/orders/{oid}/manual-status")
async def set_manual_status(oid: str, data: ManualStatusIn, request: Request, user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    if not data.reason.strip():
        raise HTTPException(400, "Begrunnelse er påkrevd ved manuell statusendring")
    allowed = {"bestilt", "delvis_levert", "levert", "manko", "restbestilling", "arkivert"}
    if data.status not in allowed:
        raise HTTPException(400, "Ugyldig status")
    order = await db.purchase_orders.find_one({"id": oid, "tenant_id": tid})
    if not order:
        raise HTTPException(404, "Fant ikke")
    entry = {"status": data.status, "at": now_iso(), "by": user.get("name"),
             "note": f"Manuell: {data.reason}", "manual": True, "from_status": order.get("status")}
    await db.purchase_orders.update_one({"id": oid, "tenant_id": tid},
                                        {"$set": {"status": data.status,
                                                  "manual_status": data.status,
                                                  "manual_status_reason": data.reason},
                                         "$push": {"status_history": entry}})
    return {"ok": True}


@api.post("/orders/{oid}/resume-auto-status")
async def resume_auto_status(oid: str, request: Request, user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    order = await db.purchase_orders.find_one({"id": oid, "tenant_id": tid})
    if not order:
        raise HTTPException(404, "Fant ikke")
    lines = order.get("lines") or []
    total_q = sum(l.get("quantity") or 0 for l in lines)
    recv_q = sum(l.get("received_quantity") or 0 for l in lines)
    auto = "bestilt"
    if total_q > 0 and recv_q >= total_q: auto = "levert"
    elif recv_q > 0: auto = "delvis_levert"
    entry = {"status": auto, "at": now_iso(), "by": user.get("name"), "note": "Gjenopptok automatisk status"}
    await db.purchase_orders.update_one({"id": oid, "tenant_id": tid},
                                        {"$set": {"status": auto, "manual_status": None, "manual_status_reason": None},
                                         "$push": {"status_history": entry}})
    return {"ok": True, "status": auto}


# ============== ORDER PDF ==============
def _build_order_pdf(order: dict, club: dict, supplier: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas as pdf_canvas
    buf = io.BytesIO()
    c = pdf_canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    y = h - 20 * mm
    c.setFont("Helvetica-Bold", 20)
    c.drawString(20 * mm, y, "BESTILLING")
    c.setFont("Helvetica", 9)
    c.drawRightString(w - 20 * mm, y, f"Innkjøpsnr: {order['po_number']}")
    y -= 6 * mm
    try:
        dt = datetime.fromisoformat(order['created_at'].replace("Z","+00:00"))
        c.drawRightString(w - 20 * mm, y, f"Dato: {dt.strftime('%d.%m.%Y')}")
    except Exception: pass
    # Club
    y_i = h - 35 * mm
    c.setFont("Helvetica-Bold", 11)
    c.drawString(20 * mm, y_i, (club or {}).get("name") or (club or {}).get("club_name") or "Klubb")
    c.setFont("Helvetica", 9)
    y_i -= 4 * mm
    for key in ("address", "postal_code", "city", "email", "phone"):
        v = (club or {}).get(key)
        if v: c.drawString(20 * mm, y_i, str(v)); y_i -= 4 * mm
    # Supplier
    y_s = h - 35 * mm
    c.setFont("Helvetica-Bold", 10)
    c.drawString(110 * mm, y_s, "Leverandør:")
    c.setFont("Helvetica", 9)
    y_s -= 4 * mm
    c.drawString(110 * mm, y_s, supplier.get("name", "")); y_s -= 4 * mm
    if supplier.get("address"): c.drawString(110 * mm, y_s, supplier["address"]); y_s -= 4 * mm
    if supplier.get("postal_code") or supplier.get("city"):
        c.drawString(110 * mm, y_s, f"{supplier.get('postal_code','')} {supplier.get('city','')}"); y_s -= 4 * mm
    if supplier.get("email"): c.drawString(110 * mm, y_s, supplier["email"]); y_s -= 4 * mm
    # Team/season
    y = h - 70 * mm
    if order.get("team_name"):
        c.setFont("Helvetica-Bold", 10)
        c.drawString(20 * mm, y, f"Lag: {order['team_name']}  ·  Årsklasse: {order.get('age_class','')}  ·  Sesong: {order.get('season','')}")
        y -= 7 * mm
    else:
        c.setFont("Helvetica-Bold", 10)
        c.drawString(20 * mm, y, "Lagerbestilling"); y -= 7 * mm
    # Table header (Varenummer FØRST)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(20 * mm, y, "Varenr")
    c.drawString(55 * mm, y, "Vare")
    c.drawString(110 * mm, y, "Str/#")
    c.drawRightString(135 * mm, y, "Antall")
    c.drawRightString(158 * mm, y, "À-pris")
    c.drawRightString(w - 20 * mm, y, "Sum")
    y -= 2 * mm; c.line(20 * mm, y, w - 20 * mm, y); y -= 5 * mm
    c.setFont("Helvetica", 9)
    total = 0
    for ln in order.get("lines") or []:
        qty = ln.get("quantity") or 0; price = ln.get("unit_price") or 0
        s = qty * price; total += s
        full_num = ln.get("full_number") or ln.get("item_number") or ""
        c.drawString(20 * mm, y, str(full_num)[:18])
        c.drawString(55 * mm, y, (ln.get("item_name") or "")[:36])
        sz = ln.get("size") or ""
        jn = ln.get("jersey_number") or ""
        c.drawString(110 * mm, y, f"{sz}{' #'+jn if jn and jn != 'X' else ''}")
        c.drawRightString(135 * mm, y, str(qty))
        c.drawRightString(158 * mm, y, f"{price:,.2f}".replace(",", " "))
        c.drawRightString(w - 20 * mm, y, f"{s:,.2f}".replace(",", " "))
        y -= 5 * mm
        if y < 50 * mm: c.showPage(); y = h - 20 * mm
    y -= 3 * mm; c.line(20 * mm, y, w - 20 * mm, y); y -= 6 * mm
    c.setFont("Helvetica-Bold", 11)
    c.drawString(110 * mm, y, "Totalt (eks. mva):")
    c.drawRightString(w - 20 * mm, y, f"NOK {total:,.2f}".replace(",", " "))
    y -= 5 * mm
    c.setFont("Helvetica", 8)
    c.drawString(110 * mm, y, "MVA 25% kommer i tillegg der det er aktuelt.")
    if order.get("notes"):
        y -= 10 * mm
        c.setFont("Helvetica-Oblique", 9)
        c.drawString(20 * mm, y, f"Notat: {order['notes'][:150]}")
    c.save()
    return buf.getvalue()


@api.get("/orders/{oid}/pdf")
async def get_order_pdf(oid: str, request: Request, user=Depends(get_current_user)):
    tid = await require_tenant(request, user)
    order = await db.purchase_orders.find_one({"id": oid, "tenant_id": tid}, {"_id": 0})
    if not order: raise HTTPException(404, "Fant ikke")
    supplier = await db.suppliers.find_one({"id": order["supplier_id"], "tenant_id": tid}, {"_id": 0}) or {}
    cust = await db.customers.find_one({"id": tid}, {"_id": 0}) or {}
    club = await db.club_settings.find_one({"tenant_id": tid}, {"_id": 0}) or {}
    if not club.get("name"):
        club = {"name": cust.get("club_name"), "email": cust.get("owner_email"),
                "phone": cust.get("owner_phone"), "address": cust.get("address", ""),
                "postal_code": cust.get("postal_code", ""), "city": cust.get("city", "")}
    try:
        pdf = _build_order_pdf(order, club, supplier)
    except Exception as e:
        logger.exception("PDF generation failed")
        raise HTTPException(500, f"PDF-feil: {e}")
    return Response(content=pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="bestilling_{order["po_number"]}.pdf"'})


# ============== UDEKKET BEHOV + RESERVATIONS ==============
@api.get("/behovsskjema/{sid}/udekket")
async def get_udekket_behov(sid: str, request: Request, user=Depends(require_perm("teams", "read"))):
    """Returnerer udekket behov for et skjema, basert på behovs-import OG tidligere bestilte mengder."""
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": sid, "tenant_id": tid}, {"_id": 0})
    if not bs: raise HTTPException(404, "Fant ikke")
    # Fetch team-needs for this team+season
    needs = await db.team_needs.find({"tenant_id": tid, "team_id": bs["team_id"], "season": bs["season"]}, {"_id": 0}).to_list(500)
    coverage = await _compute_need_coverage(tid, needs)
    # Reduce must_order by amounts already in existing non-archived orders for this behovsskjema
    prev_orders = await db.purchase_orders.find({"tenant_id": tid, "behovsskjema_id": sid,
                                                 "status": {"$nin": ["arkivert"]}}, {"_id": 0}).to_list(500)
    already = {}
    for o in prev_orders:
        for ln in o.get("lines") or []:
            k = f"{ln.get('item_id')}|{ln.get('size') or ''}|{ln.get('jersey_number') or ''}"
            already[k] = already.get(k, 0) + int(ln.get("quantity") or 0)
    # Find open stock orders (LB) that could reserve
    stock_orders = await db.purchase_orders.find({"tenant_id": tid, "is_stock_order": True,
                                                   "status": {"$in": ["bestilt", "delvis_levert"]}}, {"_id": 0}).to_list(500)
    stock_available = {}
    for o in stock_orders:
        for ln in o.get("lines") or []:
            if ln.get("reserved_for_po"): continue
            remaining = (ln.get("quantity") or 0) - (ln.get("received_quantity") or 0)
            if remaining <= 0: continue
            k = f"{ln.get('item_id')}|{ln.get('size') or ''}|{ln.get('jersey_number') or ''}"
            stock_available.setdefault(k, []).append({
                "order_id": o["id"], "line_id": ln["id"], "po_number": o["po_number"],
                "remaining": remaining,
            })
    out = []
    for l in coverage:
        k = f"{l['item_id']}|{l.get('size') or ''}|{l.get('jersey_number') or ''}"
        already_ordered = already.get(k, 0)
        remaining_to_order = max(0, l["must_order"] - already_ordered)
        out.append({**l, "already_ordered": already_ordered,
                    "remaining_to_order": remaining_to_order,
                    "stock_order_candidates": stock_available.get(k, [])})
    return {"behovsskjema": bs, "lines": out}


class ReserveStockIn(BaseModel):
    stock_order_id: str
    stock_line_id: str
    behovsskjema_id: str
    quantity: int


@api.post("/stock-reservations")
async def reserve_from_stock_order(data: ReserveStockIn, request: Request, user=Depends(require_perm("inventory", "write"))):
    tid = await require_tenant(request, user)
    bs = await db.behovsskjema.find_one({"id": data.behovsskjema_id, "tenant_id": tid}, {"_id": 0})
    if not bs: raise HTTPException(404, "Behovsskjema finnes ikke")
    order = await db.purchase_orders.find_one({"id": data.stock_order_id, "tenant_id": tid})
    if not order: raise HTTPException(404, "Lagerbestilling finnes ikke")
    lines = order.get("lines") or []
    found = next((l for l in lines if l["id"] == data.stock_line_id), None)
    if not found: raise HTTPException(404, "Linje finnes ikke")
    if found.get("reserved_for_po"):
        raise HTTPException(400, f"Linje er allerede reservert til {found['reserved_for_po']}")
    found["reserved_for_po"] = bs["po_number"]
    found["reserved_for_team"] = bs["team_name"]
    found["reserved_quantity"] = data.quantity
    await db.purchase_orders.update_one({"id": data.stock_order_id, "tenant_id": tid}, {"$set": {"lines": lines}})
    return {"ok": True, "reserved_po": bs["po_number"]}


# ============== NOTIFY LAGLEDER (SMS + email) ==============
class NotifyIn(BaseModel):
    channel: str  # email | sms | both
    message: Optional[str] = None  # editable preview; falls back to default


async def _send_sms(to: str, body: str) -> dict:
    """Returns {ok, provider_id, error}. Requires Twilio env vars."""
    if not (TWILIO_ACCOUNT_SID and TWILIO_AUTH_TOKEN and TWILIO_FROM_NUMBER):
        return {"ok": False, "error": "Twilio er ikke konfigurert. Legg til TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN og TWILIO_FROM_NUMBER i miljøvariabler."}
    try:
        from twilio.rest import Client
        c = Client(TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN)
        m = c.messages.create(body=body, from_=TWILIO_FROM_NUMBER, to=to)
        return {"ok": True, "provider_id": m.sid}
    except Exception as e:
        return {"ok": False, "error": str(e)}


@api.post("/orders/{oid}/notify-lagleder")
async def notify_lagleder(oid: str, data: NotifyIn, request: Request, user=Depends(require_perm("teams", "write"))):
    tid = await require_tenant(request, user)
    order = await db.purchase_orders.find_one({"id": oid, "tenant_id": tid}, {"_id": 0})
    if not order: raise HTTPException(404, "Fant ikke")
    if not order.get("team_id"):
        raise HTTPException(400, "Dette er en lagerbestilling – lagleder kan ikke varsles.")
    # Verify ready: all lines delivered
    lines = order.get("lines") or []
    total_q = sum(l.get("quantity") or 0 for l in lines)
    recv_q = sum(l.get("received_quantity") or 0 for l in lines)
    if total_q == 0 or recv_q < total_q:
        raise HTTPException(400, f"Bestillingen er ikke fullstendig mottatt ({recv_q}/{total_q}). Varsling er deaktivert.")
    if order.get("ready_for_pickup_notified"):
        raise HTTPException(409, "Lagleder er allerede varslet for denne bestillingen.")
    team = await db.teams.find_one({"id": order["team_id"], "tenant_id": tid}, {"_id": 0}) or {}
    cust = await db.customers.find_one({"id": tid}, {"_id": 0}) or {}
    tm = team.get("team_manager") or {}
    hc = team.get("head_coach") or {}
    recipient = tm if (tm.get("email") or tm.get("phone")) else hc
    if not (recipient.get("email") or recipient.get("phone")):
        raise HTTPException(400, "Mangler kontaktinformasjon på lagleder/hovedtrener.")
    club_name = cust.get("club_name") or "klubben"
    contact = cust.get("owner_name") or cust.get("owner_email") or ""
    default_msg = (
        f"Hei! Utstyret til {order['team_name']}, sesong {order.get('season','')}, "
        f"med innkjøpsnummer {order['po_number']}, har nå ankommet og er mottatt. "
        f"Varene er klare for henting. Ta kontakt med {contact} for å avtale henting. "
        f"OBS! Denne meldingen kan ikke besvares.\n\nHilsen {club_name}"
    )
    body = (data.message or default_msg).strip()
    log = []
    user_email = user.get("email")
    # Build attachments (order PDF) for email
    attachments = None
    try:
        supplier = await db.suppliers.find_one({"id": order["supplier_id"], "tenant_id": tid}, {"_id": 0}) or {}
        club = await db.club_settings.find_one({"tenant_id": tid}, {"_id": 0}) or {}
        if not club.get("name"):
            club = {"name": cust.get("club_name"), "email": cust.get("owner_email"),
                    "phone": cust.get("owner_phone"), "address": cust.get("address", ""),
                    "postal_code": cust.get("postal_code", ""), "city": cust.get("city", "")}
        pdf_bytes = _build_order_pdf(order, club, supplier)
        attachments = [{
            "filename": f"bestilling_{order['po_number']}.pdf",
            "content": _b64mod.b64encode(pdf_bytes).decode(),
            "content_type": "application/pdf",
        }]
    except Exception as e:
        logger.warning(f"Kunne ikke generere PDF for varsling: {e}")
    if data.channel in ("sms", "both") and recipient.get("phone"):
        r = await _send_sms(recipient["phone"], body)
        log.append({"channel": "sms", "to": recipient["phone"], "at": now_iso(), **r})
    if data.channel in ("email", "both") and recipient.get("email"):
        html = f'<pre style="font-family:Arial;font-size:13px;white-space:pre-wrap">{body}</pre>'
        r = await send_tenant_email(tid, recipient["email"],
                                    f"Utstyret for {order['team_name']} er klart",
                                    html, attachments=attachments,
                                    cc=[user_email] if user_email else None,
                                    sender_user=user)
        log.append({"channel": "email", "to": recipient["email"], "at": now_iso(), **r})
    any_ok = any(x.get("ok") for x in log)
    update = {"$push": {"notification_log": {"$each": log}}}
    if any_ok:
        update["$set"] = {"ready_for_pickup_notified": True, "notified_at": now_iso()}
    await db.purchase_orders.update_one({"id": oid, "tenant_id": tid}, update)
    return {"ok": any_ok, "log": log}


# =========== Send purchase order PDF to supplier/anyone ===========
class OrderSendIn(BaseModel):
    to: EmailStr
    cc: Optional[EmailStr] = None
    subject: Optional[str] = None
    message: Optional[str] = None


@api.post("/orders/{oid}/send")
async def send_order_email(oid: str, data: OrderSendIn, request: Request, user=Depends(require_perm("teams", "write"))):
    """Send bestillings-PDF som vedlegg til leverandør eller annen mottaker."""
    tid = await require_tenant(request, user)
    order = await db.purchase_orders.find_one({"id": oid, "tenant_id": tid}, {"_id": 0})
    if not order: raise HTTPException(404, "Fant ikke bestilling")
    supplier = await db.suppliers.find_one({"id": order["supplier_id"], "tenant_id": tid}, {"_id": 0}) or {}
    cust = await db.customers.find_one({"id": tid}, {"_id": 0}) or {}
    club = await db.club_settings.find_one({"tenant_id": tid}, {"_id": 0}) or {}
    if not club.get("name"):
        club = {"name": cust.get("club_name"), "email": cust.get("owner_email"),
                "phone": cust.get("owner_phone"), "address": cust.get("address", ""),
                "postal_code": cust.get("postal_code", ""), "city": cust.get("city", "")}
    try:
        pdf_bytes = _build_order_pdf(order, club, supplier)
    except Exception as e:
        logger.exception("PDF generation failed")
        raise HTTPException(500, f"PDF-feil: {e}")
    pdf_b64 = _b64mod.b64encode(pdf_bytes).decode()
    filename = f"bestilling_{order['po_number']}.pdf"
    club_name = club.get("name") or "klubben"
    subject = (data.subject or f"Bestilling {order['po_number']} fra {club_name}").strip()
    msg = (data.message or "").strip()
    html = (
        f'<p>Hei {supplier.get("name") or ""},</p>'
        f'<p>Vedlagt følger bestilling <strong>{order["po_number"]}</strong> fra {club_name}.</p>'
        + (f'<p>{msg}</p>' if msg else '')
        + f'<p>Vennligst bekreft mottak og forventet leveringsdato.</p>'
        f'<p>Mvh<br/>{club_name}</p>'
    )
    user_email = user.get("email")
    cc_list = []
    if data.cc: cc_list.append(data.cc)
    if user_email and user_email not in cc_list and user_email != data.to: cc_list.append(user_email)
    attachments = [{"filename": filename, "content": pdf_b64, "content_type": "application/pdf"}]
    res = await send_tenant_email(tid, data.to, subject, html,
                                   attachments=attachments, cc=cc_list or None,
                                   sender_user=user)
    if not res.get("ok"):
        raise HTTPException(502, f"E-post feilet: {res.get('error')}")
    await db.purchase_orders.update_one({"id": oid, "tenant_id": tid}, {"$set": {
        "sent_at": now_iso(), "sent_to": data.to, "last_email_id": res.get("provider_id"),
    }, "$push": {"status_history": {"status": order.get("status"), "at": now_iso(),
                                     "by": user.get("name"), "note": f"PDF sendt til {data.to}"}}})
    return {"ok": True, "sent_to": data.to, "provider_id": res.get("provider_id"), "method": res.get("method")}


@api.get("/teams/{team_id}/purchase-history")
async def team_purchase_history(team_id: str, request: Request, user=Depends(require_perm("teams", "read"))):
    """Lagshistorikk: Total innkjøpspris per år og per måned."""
    tid = await require_tenant(request, user)
    orders = await db.purchase_orders.find({"tenant_id": tid, "team_id": team_id}, {"_id": 0}).sort("created_at", -1).to_list(1000)
    years = {}
    for o in orders:
        try:
            dt = datetime.fromisoformat(o.get("created_at").replace("Z", "+00:00") if o.get("created_at") else "")
        except Exception:
            continue
        y = dt.year; m = dt.month
        total = sum((ln.get("quantity") or 0) * (ln.get("unit_price") or 0) for ln in o.get("lines", []))
        years.setdefault(y, {"year": y, "total": 0, "months": {i: 0 for i in range(1, 13)}, "orders": []})
        years[y]["total"] += total
        years[y]["months"][m] = (years[y]["months"][m] or 0) + total
        years[y]["orders"].append({
            "id": o["id"], "po_number": o["po_number"],
            "supplier_name": o.get("supplier_name"), "status": o.get("status"),
            "total": total, "created_at": o.get("created_at"),
            "season": o.get("season"),
        })
    return {"years": sorted(years.values(), key=lambda x: -x["year"]), "orders": orders}
